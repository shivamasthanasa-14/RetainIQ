import io
import re
from datetime import date, datetime, timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .store import now

COLUMNS = ['ticket_id', 'customer', 'received_at', 'subject', 'message', 'operation']
REQUIRED = COLUMNS[:-1]

def template_bytes():
    wb = Workbook()
    instructions = wb.active
    instructions.title = 'Instructions'
    for row in [
        ['RetainIQ | RelayOps customer request import'],
        ['1. Replace the example in Tickets with your own requests.'],
        ['2. Required: ticket_id, customer, received_at, subject, message.'],
        ['3. operation: new (default) or update. IDs must be unique in each workbook.'],
        ['4. received_at: ISO date/time or Excel date, interpreted as UTC if timezone absent.'],
        ['5. Updates replace intake text; received_at is immutable. Full prior values remain in history.'],
        ['6. An update invalidates prior AI processing and drafts; resolved tickets reopen.'],
        ['7. Existing escalations remain visible. Owner and follow-up are retained.'],
        ['8. Validation is all-or-nothing; correct errors, preview, then confirm import.'],
        ['9. Do not include real customer personal information in this demonstration.'],
    ]:
        instructions.append(row)
    ws = wb.create_sheet('Tickets')
    ws.append(COLUMNS)
    ws.append(['EXAMPLE-001', 'Northstar Labs', datetime(2026, 9, 29, 8), 'Invite a colleague',
               'Hi, how can I invite a teammate with editor permissions?', 'new'])
    ws['C2'].number_format = 'yyyy-mm-dd hh:mm'
    ws.auto_filter.ref = 'A1:F2'
    dv = DataValidation(type='list', formula1='"new,update"')
    ws.add_data_validation(dv)
    dv.add('F2:F1001')
    dictionary = wb.create_sheet('Data Dictionary')
    dictionary.append(['Field', 'Required', 'Accepted values / rules'])
    for row in [
        ['ticket_id', 'Yes', '1–64 letters, digits, hyphen or underscore. Stable identifier.'],
        ['customer', 'Yes', 'Company name, up to 160 characters.'],
        ['received_at', 'Yes', 'ISO date/time or Excel date; no future dates.'],
        ['subject', 'Yes', 'Up to 240 characters.'],
        ['message', 'Yes', 'Original customer message, up to 20000 characters.'],
        ['operation', 'No', 'new or update; blank means new. Updates require existing ID.'],
    ]:
        dictionary.append(row)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='14243B')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 48
    instructions.column_dimensions['A'].width = 108
    for col, width in zip('ABCDEF', [22, 26, 24, 38, 75, 15]):
        ws.column_dimensions[col].width = width
    for col, width in zip('ABC', [24, 14, 95]):
        dictionary.column_dimensions[col].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def read_workbook(data):
    if len(data) > 5_000_000:
        raise ValueError('Workbook exceeds 5 MB limit')
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if 'Tickets' not in wb.sheetnames:
            raise ValueError('Missing Tickets worksheet')
        rows = wb['Tickets'].iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if len(set(headers)) != len(headers) or not set(REQUIRED).issubset(headers):
            raise ValueError('Missing or duplicate column headers')
        if set(headers) - set(COLUMNS):
            raise ValueError('Unexpected columns; use the template')
        records = []
        for row in rows:
            if any(v is not None for v in row):
                records.append(dict(zip(headers, row)))
            if len(records) > 1000:
                raise ValueError('Maximum 1000 records per import')
        return records
    finally:
        wb.close()

def validate_records(store, records):
    errors, normalized, seen = [], [], set()
    existing = {t['ticket_id']: t for t in store.tickets()}
    for index, record in enumerate(records, 2):
        row = {}
        for field in REQUIRED:
            value = record.get(field)
            row[field] = str(value).strip() if value is not None else ''
            if not row[field]:
                errors.append(f'Row {index}: missing {field}')
            if isinstance(value, str) and value.startswith('='):
                errors.append(f'Row {index}: formulas are not accepted')
        tid = row['ticket_id']
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', tid):
            errors.append(f'Row {index}: invalid ticket_id')
        if tid in seen:
            errors.append(f'Row {index}: duplicate ticket ID {tid}')
        seen.add(tid)
        for field, limit in [('customer', 160), ('subject', 240), ('message', 20000)]:
            if len(row[field]) > limit:
                errors.append(f'Row {index}: {field} too long')
        try:
            raw = record.get('received_at')
            dt = datetime.combine(raw, datetime.min.time()) if isinstance(raw, date) and not isinstance(raw, datetime) else raw
            if not isinstance(dt, datetime):
                dt = datetime.fromisoformat(str(dt).replace('Z', '+00:00'))
            dt = dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)
            if dt > datetime.now(timezone.utc):
                raise ValueError('future date')
            row['received_at'] = dt.isoformat(timespec='seconds')
        except (ValueError, TypeError):
            errors.append(f'Row {index}: invalid or future received_at')
        op = str(record.get('operation') or 'new').strip().lower()
        row['operation'] = op
        if op not in {'new', 'update'}:
            errors.append(f'Row {index}: operation must be new or update')
        elif op == 'new' and tid in existing:
            errors.append(f'Row {index}: {tid} already exists; use update')
        elif op == 'update' and tid not in existing:
            errors.append(f'Row {index}: cannot update unknown ticket {tid}')
        normalized.append(row)
    if not records:
        errors.append('No records to import')
    return normalized, errors

def import_records(store, records, actor='Demo owner'):
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        normalized, errors = validate_records(store, records)
        if errors:
            raise ValueError('\n'.join(errors))
        for row in normalized:
            tid = row['ticket_id']
            fields = [row[c] for c in REQUIRED]
            if row['operation'] == 'new':
                db.execute('INSERT INTO tickets(ticket_id,customer,received_at,subject,message,updated_at) VALUES (?,?,?,?,?,?)', (*fields, now()))
                store.event(db, tid, 'Imported', actor, row)
                store.lifecycle_event(db, tid, 'Open', '', row['received_at'])
            else:
                before = dict(db.execute('SELECT * FROM tickets WHERE ticket_id=?', (tid,)).fetchone())
                if all(before[k] == row[k] for k in REQUIRED):
                    raise ValueError(f'{tid}: update makes no changes')
                db.execute("UPDATE drafts SET status='Superseded' WHERE ticket_id=?", (tid,))
                if row['received_at'] != before['received_at']:
                    raise ValueError(f'{tid}: received_at is immutable for updates; preserves historical reporting')
                db.execute("UPDATE tickets SET customer=?,received_at=?,subject=?,message=?,category='Unclassified',processing_status='Pending AI',summary='',ai_provider='',issue_group='',attention_reason='',missing_information='[]',result_origin='',revision=revision+1,status=CASE WHEN status='Resolved' THEN 'Open' ELSE status END,resolved_at=NULL,updated_at=? WHERE ticket_id=?", (*fields[1:], now(), tid))
                store.lifecycle_event(db, tid, 'Open' if before['status'] == 'Resolved' else before['status'], before['due_date'])
                store.event(db, tid, 'Intake updated', actor, {'before': before, 'after': row})
    return len(normalized)
