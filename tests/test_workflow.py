import io
import json
from datetime import date
from pathlib import Path
import pytest
from openpyxl import load_workbook
from retainiq.ai import AIUnavailable, validate_classification, validate_draft
from retainiq.constants import OWNERS
from retainiq.demo import initialize
from retainiq.dispatch import dispatch
from retainiq.imports import template_bytes, read_workbook, import_records, validate_records
from retainiq.knowledge import Knowledge
from retainiq.store import Store
from retainiq.workflow import Workflow, metrics

class FakeAI:
    label = 'TEST DOUBLE — not live AI'
    calls = 0
    unsupported = False
    fail = False
    def classify(self, ticket):
        self.calls += 1
        if self.fail:
            raise AIUnavailable('Test outage')
        return {'category': 'Integrations' if ticket['subject'] == 'outage' else 'Account Access',
                'summary': ticket['subject'], 'needs_specialist': ticket['subject'] == 'outage',
                'reason': 'Production failure' if ticket['subject'] == 'outage' else 'Standard request', 'missing_information': []}
    def draft(self, ticket, docs):
        return {'supported': not self.unsupported, 'body': 'An Admin can invite users from Settings > Members > Invite.',
                'source_ids': [docs[0]['id']], 'missing_information': []}

@pytest.fixture
def setup(tmp_path):
    store = Store(tmp_path / 'test.db')
    initialize(tmp_path, store)
    with store.connect() as db:
        db.execute('DELETE FROM history')
        db.execute('DELETE FROM lifecycle')
        db.execute('DELETE FROM tickets')
    ai = FakeAI()
    return store, ai, Workflow(store, ai, Knowledge(tmp_path / 'knowledge_base'))

def record(tid='TEST-1', subject='Invite colleague', message='How can an Admin invite a teammate as an Editor?'):
    return {'ticket_id': tid, 'customer': 'Example Co', 'received_at': '2026-01-01', 'subject': subject, 'message': message}

def test_template():
    data = template_bytes()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ['Instructions', 'Tickets', 'Data Dictionary']
    assert len(read_workbook(data)) == 1
    assert wb['Tickets'].data_validations.count == 1

def test_import_persistence(setup):
    store, _, _ = setup
    assert import_records(store, [record()]) == 1
    assert Store(store.path).ticket('TEST-1')['message'] == record()['message']
    assert store.history('TEST-1')[0]['action'] == 'Imported'

@pytest.mark.parametrize('bad', [dict(record(), message=''), dict(record(), received_at='invalid'), dict(record(), operation='delete'), dict(record(), customer='=1+1'), dict(record(), received_at='2999-01-01'), dict(record(), ticket_id='bad id')])
def test_invalid_atomic_import(setup, bad):
    store, _, _ = setup
    with pytest.raises(ValueError):
        import_records(store, [record('GOOD'), bad])
    assert not store.tickets()

def test_duplicates(setup):
    store, _, _ = setup
    with pytest.raises(ValueError, match='duplicate'):
        import_records(store, [record(), record()])
    import_records(store, [record()])
    with pytest.raises(ValueError, match='already exists'):
        import_records(store, [record()])
    with pytest.raises(ValueError, match='unknown'):
        import_records(store, [dict(record('NEW'), operation='update')])

@pytest.mark.parametrize('changes', [{'category':'unknown'}, {'needs_specialist':'yes'}, {'summary':''}, {'missing_information':'text'}])
def test_classification_schema(changes):
    valid = FakeAI().classify(record())
    with pytest.raises(ValueError):
        validate_classification({**valid, **changes})

def test_retrieval(setup):
    _, _, workflow = setup
    hits = workflow.knowledge.retrieve('invite colleague editor role account permissions')
    assert hits[0]['id'] == '02-account-permissions'
    assert not workflow.knowledge.retrieve('xylophone quasar unicorn')

def test_draft_edit_approve(setup):
    store, _, workflow = setup
    import_records(store, [record()])
    workflow.process('TEST-1')
    did = workflow.prepare_draft('TEST-1')
    assert workflow.prepare_draft('TEST-1') == did
    workflow.edit_draft(did, 'Reviewed response', 'Reviewer')
    assert store.draft('TEST-1')[0]['original_body'] != 'Reviewed response'
    with pytest.raises(ValueError):
        workflow.approve(did, '')
    workflow.approve(did, 'Reviewer')
    draft = store.draft('TEST-1')[0]
    assert draft['approved_at'] and draft['approved_by'] == 'Reviewer'
    assert store.ticket('TEST-1')['status'] == 'Open'
    with pytest.raises(ValueError):
        workflow.approve(did, 'Reviewer')
    with pytest.raises(ValueError):
        workflow.edit_draft(did, 'new edit', 'Reviewer')

def test_unsupported_answer(setup):
    store, ai, workflow = setup
    ai.unsupported = True
    import_records(store, [record()])
    workflow.process('TEST-1')
    assert workflow.prepare_draft('TEST-1') is None
    assert store.ticket('TEST-1')['status'] == 'Escalated'
    assert not store.draft('TEST-1')

def test_no_docs(setup):
    store, _, workflow = setup
    import_records(store, [record(message='xylophone quasar unicorn', subject='quasar')])
    workflow.process('TEST-1')
    assert workflow.prepare_draft('TEST-1') is None
    assert store.ticket('TEST-1')['status'] == 'Escalated'

def test_draft_source_schema():
    with pytest.raises(ValueError):
        validate_draft({'supported':True,'body':'made up','source_ids':['absent'],'missing_information':[]}, [])

def test_batch_duplicate_prevention(setup):
    store, ai, workflow = setup
    import_records(store, [record(), record('TEST-2', 'outage')])
    results = workflow.batch(['TEST-1', 'TEST-2', 'TEST-1'])
    assert len(results) == 2 and ai.calls == 2
    assert workflow.process('TEST-1').startswith('Skipped')
    assert ai.calls == 2
    assert store.ticket('TEST-2')['status'] == 'Escalated'
    assert store.ticket('TEST-2')['owner'] == OWNERS[2]
    assert len(store.query('SELECT * FROM dispatches')) == 1
    with pytest.raises(ValueError):
        workflow.escalate('TEST-2', OWNERS[2], 'duplicate', 'Reviewer')

def test_failure_retry(setup):
    store, ai, workflow = setup
    import_records(store, [record()])
    ai.fail = True
    assert workflow.process('TEST-1').startswith('AI Failed')
    assert store.ticket('TEST-1')['category'] == 'Unclassified'
    ai.fail = False
    assert workflow.process('TEST-1') == 'Processed'
    assert len(store.history('TEST-1')) == 3

def test_transitions(setup):
    store, _, workflow = setup
    import_records(store, [record()])
    workflow.follow_up('TEST-1', OWNERS[0], 'Check result', '2026-01-02', 'In Progress', '', 'Reviewer')
    with pytest.raises(ValueError, match='notes'):
        workflow.follow_up('TEST-1', OWNERS[0], '', '', 'Resolved', '', 'Reviewer')
    workflow.follow_up('TEST-1', OWNERS[0], '', '', 'Resolved', 'Customer confirmed', 'Reviewer')
    assert store.ticket('TEST-1')['resolved_at']
    resolution_time=store.ticket('TEST-1')['resolved_at']
    workflow.follow_up('TEST-1', OWNERS[0], '', '', 'Resolved', 'Notes clarified', 'Reviewer')
    assert store.ticket('TEST-1')['resolved_at']==resolution_time
    with pytest.raises(ValueError, match='transition'):
        workflow.follow_up('TEST-1', OWNERS[0], 'Check', '', 'In Progress', '', 'Reviewer')
    workflow.follow_up('TEST-1', OWNERS[0], 'Reopened by customer', '', 'Open', '', 'Reviewer')
    assert store.ticket('TEST-1')['resolved_at'] is None

def test_update_preserves_history_invalidates_draft(setup):
    store, _, workflow = setup
    import_records(store, [record()])
    workflow.process('TEST-1')
    did = workflow.prepare_draft('TEST-1')
    import_records(store, [dict(record(), message='New original question', operation='update')])
    assert store.ticket('TEST-1')['revision'] == 2
    assert store.ticket('TEST-1')['processing_status'] == 'Pending AI'
    assert store.draft('TEST-1')[0]['status'] == 'Superseded'
    assert 'How can an Admin' in store.history('TEST-1')[0]['detail']
    with pytest.raises(ValueError):
        workflow.approve(did, 'Reviewer')

def test_dashboard(setup):
    store, _, workflow = setup
    import_records(store, [record(), record('TEST-2')])
    workflow.follow_up('TEST-1', OWNERS[0], 'Check', '2026-01-01', 'In Progress', '', 'Reviewer')
    workflow.follow_up('TEST-2', OWNERS[0], '', '', 'Resolved', 'Confirmed', 'Reviewer')
    m = metrics(store.tickets(), date(2026, 1, 3))
    assert m == {'New Requests':1,'Open Requests':1,'Resolved Requests':1,'Requests Needing Review':0,'Overdue Requests':1}

def test_dispatch_failure_retry_idempotency(setup, monkeypatch):
    import requests
    store, _, workflow = setup
    import_records(store, [record()])
    workflow.escalate('TEST-1', OWNERS[0], 'Review', 'Reviewer')
    def fail(*args, **kwargs):
        raise requests.ConnectionError('offline')
    monkeypatch.setattr(requests, 'post', fail)
    assert dispatch(store, 1) == 'Failed'
    class Response:
        def raise_for_status(self): pass
        def json(self): return {'received':True,'dispatch_id':1}
    monkeypatch.setattr(requests, 'post', lambda *args, **kwargs: Response())
    assert dispatch(store, 1) == 'Delivered'
    assert dispatch(store, 1) == 'Already delivered'
    assert store.query('SELECT attempts FROM dispatches')[0]['attempts'] == 2
    with pytest.raises(ValueError):
        dispatch(store, 1, 'https://external.example/receive')
